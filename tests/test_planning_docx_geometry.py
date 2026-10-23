"""Geometry of the supplied landscape unit reference, including minimal drafts."""
import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from docx_builder_planning import build_docx_from_planning_v2, TABLE_WIDTH
from models.planning_document import PlanningContainerV2


def assert_geometry(doc):
    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width.twips == 16839
    assert section.page_height.twips == 11907
    assert section.left_margin.twips == 992
    assert section.right_margin.twips == 680
    assert section.top_margin.twips >= 1080
    assert len(doc.tables) == 8
    assert [len(t.columns) for t in doc.tables] == [11, 1, 2, 3, 3, 3, 7, 2]
    assert len(doc.tables[0].rows) == 3
    for i, table in enumerate(doc.tables):
        widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid]
        assert sum(widths) == TABLE_WIDTH
        assert sum(widths) <= section.page_width.twips - section.left_margin.twips - section.right_margin.twips
        if i >= 2:
            assert table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
    assert not doc.tables[6]._tbl.xpath('.//w:cantSplit'), 'Long session rows must be able to flow across pages'
    assert section.header.tables, 'Reference masthead must use stable inline layout'
    assert not section.header._element.xpath('.//wp:anchor'), 'Floating artwork must not displace body text'


def main():
    payload = json.loads((ROOT/'data/pedagogy/fixtures/secondary_math_project_unit.v2.json').read_text(encoding='utf-8'))
    assert_geometry(Document(build_docx_from_planning_v2(PlanningContainerV2(**payload))))
    minimal = PlanningContainerV2(id='minimal-unit', identity={'planningType':'unit'})
    assert_geometry(Document(build_docx_from_planning_v2(minimal)))
    print('test_planning_docx_geometry.py: reference layout, widths, repeat headers and draft OK')


if __name__ == '__main__':
    main()
