"""Planning export preserves content in the eight reference tables."""
import copy
import json
from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'backend'))
from docx_builder_planning import build_docx_from_planning_v2
from models.planning_document import PlanningContainerV2


def text(table):
    return '\n'.join(cell.text for row in table.rows for cell in row.cells)


def main():
    payload=json.loads((ROOT/'data/pedagogy/fixtures/secondary_math_project_unit.v2.json').read_text(encoding='utf-8'))
    payload['identity']['planningType']='unit'
    payload['administrativeContext'].update(director='Director prueba',coordinator='Coordinador prueba',sections=['A','B'])
    doc=Document(build_docx_from_planning_v2(PlanningContainerV2(**payload)))
    headings='\n'.join(p.text for p in doc.paragraphs)
    assert 'UNIDAD DIDÁCTICA DE APRENDIZAJE' in headings
    assert 'PROPÓSITO Y PRODUCTO FINAL' in headings
    assert 'ESTÁNDAR DE APRENDIZAJE' in headings
    assert 'MATRIZ DE EVALUACIÓN' not in headings
    assert len(doc.tables)==8
    for value in ['Director prueba','Coordinador prueba','A, B',payload['administrativeContext']['institution']]:
        assert value in text(doc.tables[0])
    assert payload['significantSituation']['context'] in text(doc.tables[1])
    assert payload['purpose']['summary'] in text(doc.tables[2])
    assert payload['finalProduct']['title'] in text(doc.tables[2])
    assert payload['curriculumMap'][0]['standard']['description'] in text(doc.tables[3])
    sequence=doc.tables[6]
    assert len(sequence.rows)==len(payload['sequence'])+1
    for row,item in zip(sequence.rows[1:],sorted(payload['sequence'],key=lambda s:s['index'])):
        assert item['title'] in row.cells[2].text
        linked=[entry for entry in payload['curriculumMap'] if entry['id'] in item['curriculumMapRefs']]
        for entry in linked:
            for criterion in entry['criteria']:
                if criterion['id'] in item['criterionRefs']:
                    assert criterion['description'] in row.cells[3].text
    # A second area retains its own standard and capabilities.
    multi=copy.deepcopy(payload)
    entry=copy.deepcopy(multi['curriculumMap'][0])
    entry['id']='science-map';entry['area']={'id':'science','officialName':'Ciencia y Tecnología'}
    entry['competency']={'id':'science-inquiry','officialName':'Indaga mediante métodos científicos'}
    entry['standard']['description']='Estándar de ciencias de prueba'
    multi['curriculumMap'].append(entry)
    multiple=Document(build_docx_from_planning_v2(PlanningContainerV2(**multi)))
    assert 'Ciencia y Tecnología' in text(multiple.tables[0])
    assert 'Estándar de ciencias de prueba' in text(multiple.tables[3])
    # Missing values stay blank rather than inventing a school, teacher or dates.
    minimal=Document(build_docx_from_planning_v2(PlanningContainerV2(id='draft-unit',identity={})))
    assert len(minimal.tables)==8
    assert 'Agropecuario JEC' not in text(minimal.tables[0])
    print('test_planning_docx_builder.py: content, ordered sequence, multiple areas and draft OK')


if __name__=='__main__':
    main()
