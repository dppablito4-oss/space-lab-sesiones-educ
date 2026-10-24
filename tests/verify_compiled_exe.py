import json
import copy
import io
import socket
import subprocess
import sys
import time
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

import docx
import httpx
import pymupdf
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend.version import ENGINE_VERSION  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
# CI keeps the generic release filename; a local versioned build can be supplied.
EXE_PATH = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "backend" / "dist" / "pablitopyhost.exe"
TOKEN_PATH = EXE_PATH.parent / "connection_token.txt"
FIXTURE_PATH = ROOT / "data" / "pedagogy" / "fixtures" / "secondary_math_project_unit.v2.json"
OUTPUT_DOCX = ROOT / "tests" / "fixtures" / "test_compiled_exe_output.docx"


def main():
    print(f"[1/5] Verificando existencia de {EXE_PATH.name}...")
    assert EXE_PATH.exists(), f"No se encontró el ejecutable en {EXE_PATH}"
    print(f"  [OK] Binario encontrado ({EXE_PATH.stat().st_size:,} bytes)")

    with socket.socket() as probe:
        probe.settimeout(0.5)
        if probe.connect_ex(('127.0.0.1', 8000)) == 0:
            raise RuntimeError('El puerto 8000 está ocupado; cierra el motor antes de verificar el ejecutable nuevo.')

    print(f"\n[2/5] Levantando proceso ejecutable {EXE_PATH.name}...")
    startup = subprocess.STARTUPINFO() if sys.platform == 'win32' else None
    if startup:
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = 0
    process = subprocess.Popen(
        [str(EXE_PATH)],
        cwd=str(EXE_PATH.parent),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        startupinfo=startup,
        creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == 'win32' else 0,
    )

    try:
        online = False
        start_time = time.time()
        while time.time() - start_time < 45:
            try:
                res = httpx.get("http://127.0.0.1:8000/", timeout=1.0)
                if res.status_code == 200 and res.json().get("status") == "Online":
                    online = True
                    server_data = res.json()
                    break
            except Exception:
                time.sleep(0.5)

        assert online, "El ejecutable no respondió en http://127.0.0.1:8000 en 45 segundos"
        print(f"  [OK] Servidor local en línea: {server_data}")
        assert server_data.get("version") == ENGINE_VERSION

        print("\n[3/5] Leyendo token criptográfico generado por el ejecutable...")
        assert TOKEN_PATH.exists(), f"No se generó {TOKEN_PATH}"
        token = TOKEN_PATH.read_text(encoding="utf-8").strip()
        assert len(token) == 64, "El token debe tener 64 caracteres hexadecimales"
        print("  [OK] Token de conexión presente (valor oculto)")

        print("\n[4/5] Exportando unidad horizontal según referencia (8 tablas)...")
        unit_payload = json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))
        unit_payload["token"] = token

        export_res = httpx.post(
            "http://127.0.0.1:8000/exportar-docx-json",
            json=unit_payload,
            headers={
                "Origin": "https://sesiones.sypablitodp.site",
                "Content-Type": "application/json",
            },
            timeout=15.0,
        )

        assert export_res.status_code == 200, f"Error {export_res.status_code}: {export_res.text}"
        assert export_res.content[:2] == b"PK", "El contenido no es un zip/docx válido"
        print(f"  [OK] Documento Word devuelto con éxito ({len(export_res.content):,} bytes)")

        OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
        OUTPUT_DOCX.write_bytes(export_res.content)

        print("\n[5/5] Auditando estructura interna del DOCX generado por el EXE...")
        doc = docx.Document(str(OUTPUT_DOCX))
        print(f"  - Total de tablas en el documento: {len(doc.tables)}")
        assert len(doc.tables) == 8, f"Se esperaban 8 tablas, encontradas: {len(doc.tables)}"
        assert doc.sections[0].page_width > doc.sections[0].page_height

        # Verificar anchos de tabla
        for i, table in enumerate(doc.tables):
            col_widths = [cell.width for cell in table.rows[0].cells]
            total_width = sum(col_widths)
            assert total_width > 0
            print(f"    * Tabla {i+1:02d}: {len(table.rows)} filas x {len(table.columns)} cols")

        print("\nVerificando PDF real desde SessionDocument v1...")
        pdf_payload = json.loads((ROOT / 'tests/fixtures/secundaria-matematica-polya.v1.json').read_text(encoding='utf-8'))
        pdf_payload['momentos']['desarrollo']['procesos'][0]['contenido']['value'] += (
            r'<p>MATH-AUDIT-145: calculamos \(\sqrt[3]{8}\), $\sin(\theta)$ y $\int_0^1 x^2 dx$.</p>'
        )
        pdf_payload['token'] = token
        pdf_response = httpx.post(
            'http://127.0.0.1:8000/exportar-pdf-json', json=pdf_payload,
            headers={'Origin': 'https://sesiones.sypablitodp.site'}, timeout=90,
        )
        assert pdf_response.status_code == 200, f'Exportación PDF devolvió HTTP {pdf_response.status_code}'
        assert pdf_response.content.startswith(b'%PDF'), 'La respuesta debe ser un PDF válido'
        pdf_artifact = ROOT / 'artifacts/math-review/compiled-session.pdf'
        pdf_artifact.parent.mkdir(parents=True, exist_ok=True)
        pdf_artifact.write_bytes(pdf_response.content)
        with pymupdf.open(stream=pdf_response.content, filetype='pdf') as exported_pdf:
            assert 'MATH-AUDIT-145' in ''.join(page.get_text() for page in exported_pdf), 'El PDF omitió el proceso con fórmulas'
        print(f'  [OK] PDF canónico generado ({len(pdf_response.content):,} bytes)')

        print('\nVerificando planificación y sesiones en un único Word...')
        linked_sessions = []
        for item in unit_payload['sequence']:
            if item['type'] != 'session':
                continue
            session = copy.deepcopy(pdf_payload)
            session.pop('token', None)
            session['id'] = 'compiled-' + item['id']
            session['planning'] = {'planningContainerId': unit_payload['id'], 'sequenceItemId': item['id']}
            item['linkedDocumentRef'] = {'id': session['id']}
            session['metadata']['numeroSesion'] = str(item['index'])
            linked_sessions.append(session)
        unit_payload['linkedSessions'] = linked_sessions
        bundled = httpx.post('http://127.0.0.1:8000/exportar-docx-json', json=unit_payload,
                             headers={'Origin': 'https://sesiones.sypablitodp.site'}, timeout=90)
        assert bundled.status_code == 200, f'Bundle devolvió HTTP {bundled.status_code}'
        bundled_doc = docx.Document(io.BytesIO(bundled.content))
        assert sum(p.text.startswith('SESIÓN DE APRENDIZAJE') for p in bundled_doc.paragraphs) == len(linked_sessions)
        assert len(bundled_doc._element.xpath('.//m:rad/m:deg')) >= len(linked_sessions)
        assert len(bundled_doc._element.xpath('.//m:nary')) >= len(linked_sessions)
        print(f'  [OK] Word completo generado ({len(bundled.content):,} bytes)')

        print("\n============================================================")
        print(">>> CERTIFICACIÓN DEL EJECUTABLE PABLITOPYHOST.EXE EXITOSA <<<")
        print("============================================================")

    finally:
        print(f"\nCerrando proceso {EXE_PATH.name}...")
        if sys.platform == 'win32':
            subprocess.run(['taskkill.exe', '/PID', str(process.pid), '/T', '/F'],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        else:
            process.terminate()
        try:
            process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            process.kill()
        print("Proceso finalizado limpiamente.")


if __name__ == "__main__":
    main()
