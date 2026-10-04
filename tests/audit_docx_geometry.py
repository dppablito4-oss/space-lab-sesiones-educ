"""Auditoría Técnica de Estructura y Márgenes - Documento Word (.docx) Planning MINEDU.

Verifica rigurosamente:
- A4 (8.27 in x 11.69 in)
- Márgenes de 0.75 in en los 4 bordes
- Exactamente 12 tablas oficiales
- 10,490 twips exactos en todas las tablas
- Sangría de tabla w:tblInd = -289 twips
- Protección de quiebre de fila w:cantSplit al 100%
- Encabezados repetibles w:tblHeader en tablas extensas

Ejecutar: python tests/audit_docx_geometry.py [ruta_al_docx_opcional]
"""
import json
import sys
from pathlib import Path
import docx
from docx.oxml.ns import qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

# Determinar archivo a auditar
file_path = None
if len(sys.argv) > 1 and Path(sys.argv[1]).exists():
    file_path = Path(sys.argv[1])
else:
    files = sorted(Path("tests/fixtures").glob("unidad_3sec_ept_manual*.docx"), key=lambda f: f.stat().st_mtime, reverse=True)
    if files:
        file_path = files[0]

if file_path and file_path.exists():
    print(f"Abriendo archivo existente: {file_path} ({file_path.stat().st_size:,} bytes)")
    doc = docx.Document(str(file_path))
else:
    print("No se encontró archivo previo en disco. Generando documento en memoria desde fixture canónico...")
    from models.planning_document import PlanningContainerV2
    from docx_builder_planning import build_docx_from_planning_v2

    fixture_path = ROOT / "data" / "pedagogy" / "fixtures" / "secondary_math_project_unit.v2.json"
    data = json.loads(fixture_path.read_text(encoding="utf-8"))
    container = PlanningContainerV2(**data)
    stream = build_docx_from_planning_v2(container)
    doc = docx.Document(stream)
    file_path = Path("memoria:fixture_secondary_math_project_unit.docx")

print("========================================================")
print("AUDITORIA TECNICA DE ESTRUCTURA Y MARGENES - DOCUMENTO WORD")
print(f"Archivo auditado: {file_path}")
print("========================================================")

# 1. Configuración de página y márgenes A4
sec = doc.sections[0]
print("\n1. CONFIGURACION DE PAGINA Y AREA DE IMPRESION:")
print(f"   - Formato de Papel: A4")
print(f"   - Ancho Total de Hoja: {sec.page_width.inches:.3f} in ({sec.page_width.twips} twips)")
print(f"   - Alto Total de Hoja:  {sec.page_height.inches:.3f} in ({sec.page_height.twips} twips)")
print(f"   - Margen Superior:     {sec.top_margin.inches:.3f} in ({sec.top_margin.twips} twips)")
print(f"   - Margen Inferior:     {sec.bottom_margin.inches:.3f} in ({sec.bottom_margin.twips} twips)")
print(f"   - Margen Izquierdo:    {sec.left_margin.inches:.3f} in ({sec.left_margin.twips} twips)")
print(f"   - Margen Derecho:      {sec.right_margin.inches:.3f} in ({sec.right_margin.twips} twips)")

assert abs(sec.page_width.inches - 8.27) < 0.05, f"Ancho de página no cumple con A4 (8.27 in): {sec.page_width.inches}"
assert abs(sec.page_height.inches - 11.69) < 0.05, f"Alto de página no cumple con A4 (11.69 in): {sec.page_height.inches}"
assert abs(sec.top_margin.inches - 0.75) < 0.05, f"Margen superior no cumple con 0.75 in: {sec.top_margin.inches}"
assert abs(sec.bottom_margin.inches - 0.75) < 0.05, f"Margen inferior no cumple con 0.75 in: {sec.bottom_margin.inches}"
assert abs(sec.left_margin.inches - 0.75) < 0.05, f"Margen izquierdo no cumple con 0.75 in: {sec.left_margin.inches}"
assert abs(sec.right_margin.inches - 0.75) < 0.05, f"Margen derecho no cumple con 0.75 in: {sec.right_margin.inches}"

# 2. Conteo estricto de tablas: exactamente 12 tablas
print(f"\n2. AUDITORIA DE TODAS LAS TABLAS ({len(doc.tables)} tablas oficiales):")
assert len(doc.tables) == 12, f"Se esperaban exactamente 12 tablas, encontradas {len(doc.tables)}"

all_10490 = True
all_cant_split = True
all_ind_correct = True

table_names = [
    "I. Ficha Técnica / Datos Informativos",
    "II. Situación Significativa y Reto",
    "II. Tabla Dual: Propósito vs Producto Final",
    "III. Estándares de Aprendizaje (Ciclo)",
    "III. Matriz Curricular (Criterios y Desempeños)",
    "IV. Competencias Transversales CNEB",
    "IV. Enfoques Transversales y Valores",
    "V. Secuencia Didáctica Semanal (7 Columnas)",
    "VI. Producto Final Integrador",
    "VII. Materiales de Aula y Recursos MINEDU",
    "VIII. Orientaciones de Evaluación Formativa",
    "IX. Firmas de Responsabilidad Pedagógica"
]

repeat_header_indices = {2, 3, 4, 5, 6, 7, 8, 9, 10}

for i, tbl in enumerate(doc.tables):
    tblGrid = tbl._tbl.find(qn('w:tblGrid'))
    col_widths = [int(gc.attrib.get(qn('w:w'), 0)) for gc in tblGrid.findall(qn('w:gridCol'))] if tblGrid is not None else []
    total_w = sum(col_widths)
    if total_w != 10490:
        all_10490 = False

    first_tr = tbl.rows[0]._tr.get_or_add_trPr()
    has_header = first_tr.find(qn('w:tblHeader')) is not None
    has_split = all(r._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None for r in tbl.rows)
    if not has_split:
        all_cant_split = False

    ind_el = tbl._tbl.tblPr.find(qn('w:tblInd'))
    ind_val = ind_el.attrib.get(qn('w:w')) if ind_el is not None else '0'
    if int(ind_val) != -289:
        all_ind_correct = False

    name = table_names[i] if i < len(table_names) else f"Tabla {i}"
    print(f"   [{i:02d}] {name:<46} | {len(tbl.rows)}x{len(col_widths)} cols | sum={total_w} twips | ind={ind_val} twips | repeatHeader={has_header} | noRowSplit={has_split}")

    assert total_w == 10490, f"Tabla {i} [{name}] no cumple con 10,490 twips (tiene {total_w})"
    assert int(ind_val) == -289, f"Tabla {i} [{name}] no tiene sangría -289 twips (tiene {ind_val})"
    assert has_split, f"Tabla {i} [{name}] contiene filas sin w:cantSplit"
    if i in repeat_header_indices:
        assert has_header, f"Tabla {i} [{name}] debe tener w:tblHeader"

print("\n3. EVALUACION DE CONFORMIDAD GEOMETRICA:")
print(f"   [OK] Cuadricula oficial 10,490 twips en 100% de tablas: {all_10490}")
print(f"   [OK] Sangria institucional (-289 twips / -0.2 in) aplicada: {all_ind_correct}")
print(f"   [OK] Renglones protegidos contra quiebre de pagina: {all_cant_split}")
print(f"   [OK] Formato A4 y margenes institucionales de 0.75 in: True")
print("========================================================")
print(">>> AUDITORÍA DE GEOMETRÍA FINALIZADA CON ÉXITO: 100% CERTIFICADO <<<")

assert all_10490 is True, "No todas las tablas cumplen con 10,490 twips"
assert all_cant_split is True, "No todas las filas están protegidas con cantSplit"
assert all_ind_correct is True, "No todas las tablas tienen sangría de -289 twips"
