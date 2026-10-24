"""Generate representative math exports; --word checks them in installed Word.

Usage: python scripts/audit_word_math_exports.py [--word]
Results are local, ignored artifacts. Word runs hidden and closes in finally.
"""
import argparse
import json
from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from word_math import append_mixed_text  # noqa: E402
from docx_builder_v1 import build_docx_from_v1  # noqa: E402
from docx_builder_planning import build_docx_from_planning_v2  # noqa: E402
from models.session_document import SessionDocumentV1  # noqa: E402
from models.planning_document import PlanningContainerV2  # noqa: E402

FORMULAS = [
    ('Fracciones', r'$\frac{3}{4}+\frac{1}{2}=\frac{5}{4}$'),
    ('Raíz cúbica', r'$\sqrt[3]{8}=2$'),
    ('Fórmula cuadrática', r'$x=\frac{-b\pm\sqrt{b^2-4ac}}{2a}$'),
    ('Potencias e índices', r'$x_{1}^{2}+x_{2}^{2}=25$'),
    ('Argumentos separados por espacios', r'$\frac {1} {2}+x ^ {2}$'),
    ('Trigonometría', r'$\sin^2(\theta)+\cos^2(\theta)=1$'),
    ('Logaritmos', r'$\log_{2}(8)=3$'),
    ('Unidades', r'$A=12\mathrm{cm}^{2}$'),
    ('Vectores', r'$\vec{v}=3\vec{i}+4\vec{j}$'),
    ('Delimitadores', r'$\left(\frac{a}{b}\right)^2$'),
    ('Sumatoria', r'$\sum_{i=1}^{5} i=15$'),
    ('Integral', r'$\int_0^1 x^2 dx=\frac{1}{3}$'),
    ('Delimitador invisible', r'$\left. x^2\right|_0^1=1$'),
    ('Paréntesis LaTeX', r'\(\frac{2}{3}+\sqrt{9}\)'),
    ('Bloque LaTeX', r'\[\frac{1}{2}=0.5\]'),
]


def main():
    parser = argparse.ArgumentParser(__doc__)
    parser.add_argument('--word', action='store_true')
    args = parser.parse_args()
    output = ROOT / 'artifacts/math-review'
    output.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading('Fórmulas editables: revisión de exportación', 0)
    for label, formula in FORMULAS:
        doc.add_paragraph(label, style='Heading 2')
        routes = append_mixed_text(doc.add_paragraph(), formula, font_size=12)
        assert routes == ['omml'], (label, routes)
    examples = output / 'Formulas_OMML.docx'
    doc.save(examples)

    session = json.loads((ROOT / 'tests/fixtures/secundaria-matematica-polya.v1.json').read_text(encoding='utf-8'))
    session['metadata']['titulo'] = r'Aplicamos $\sqrt[3]{8}$ y fracciones'
    session['momentos']['desarrollo']['procesos'][0]['contenido']['value'] = ''.join(
        '<p>' + label + ': ' + formula + '</p>' for label, formula in FORMULAS)
    session['fichaTrabajo'] = {'titulo': r'Calculamos $\sqrt{x}$',
        'indicaciones': r'Resuelve \(\frac{1}{2}+\frac{1}{3}\).',
        'actividades': r'Comprueba $\sin^2(\theta)+\cos^2(\theta)=1$.'}
    session['listaCotejo'] = {'criterios': [r'Calcula $\frac{3}{4}$ correctamente'], 'alumnos': []}
    session_path = output / 'Sesion_con_formulas.docx'
    session_path.write_bytes(build_docx_from_v1(SessionDocumentV1(**session)).getvalue())

    plan = json.loads((ROOT / 'data/pedagogy/fixtures/secondary_math_project_unit.v2.json').read_text(encoding='utf-8'))
    plan['identity']['title'] = r'Proyecto: investigamos $\sqrt[3]{8}$ y $\frac{1}{2}$'
    plan['purpose']['summary'] = r'Aplicar \(\frac{1}{2}\) y $\sqrt[3]{8}$.'
    plan['sequence'][0]['knowledge'] = [formula for _, formula in FORMULAS[:4]]
    plan_path = output / 'Unidad_con_formulas.docx'
    plan_path.write_bytes(build_docx_from_planning_v2(PlanningContainerV2(**plan)).getvalue())

    report = {}
    app = None
    try:
        if args.word:
            import win32com.client
            app = win32com.client.DispatchEx('Word.Application')
            app.Visible = False
            app.DisplayAlerts = 0
        for path in (examples, session_path, plan_path):
            expected = len(Document(path)._element.xpath('.//m:oMath'))
            result = {'xml_equations': expected}
            if app:
                opened = app.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False, OpenAndRepair=False)
                try:
                    result['word_equations'] = opened.OMaths.Count
                    assert result['word_equations'] == expected, (path.name, result)
                    opened.Repaginate()
                    result['pages'] = opened.ComputeStatistics(2)
                    opened.ExportAsFixedFormat(str(path.with_suffix('.pdf')), 17)
                    normalized = path.with_name(path.stem + '_Word.docx')
                    opened.SaveAs2(str(normalized), FileFormat=16, AddToRecentFiles=False)
                finally:
                    opened.Close(0)
                result['saved_equations'] = len(Document(normalized)._element.xpath('.//m:oMath'))
                assert result['saved_equations'] == expected, (path.name, result)
            report[path.name] = result
    finally:
        if app:
            app.Quit()
    (output / 'audit.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
