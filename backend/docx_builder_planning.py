"""Landscape planning export following the supplied Monte Alegre unit layout."""
from html import escape
from io import BytesIO

from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from document_template import new_document_from_template
from docx_builder import (add_table_borders_black, append_html_to_cell_or_paragraph,
                          set_cell_background, set_cell_margins,
                          set_table_col_widths_and_indent)
from models.planning_document import PlanningContainerV2
from word_math import append_mixed_text

PAGE_WIDTH, PAGE_HEIGHT, TABLE_WIDTH = 16839, 11907, 15167
COLORS = ('0070C0', 'FF0000', '00A000', '7030A0')


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (list, tuple)):
        return '\n'.join(filter(None, (_text(item) for item in value)))
    if isinstance(value, dict):
        text = _text(value.get('description') or value.get('citation') or value.get('title') or value.get('name') or value.get('label') or '')
        url = _text(value.get('url'))
        return text + ('\n' + url if url and url not in text else '')
    return str(value).strip()


def build_docx_from_planning_v2(container: PlanningContainerV2) -> BytesIO:
    doc = new_document_from_template('planning')
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Twips(PAGE_WIDTH), Twips(PAGE_HEIGHT)
        section.top_margin, section.bottom_margin = Twips(1080), Twips(720)
        section.left_margin, section.right_margin = Twips(992), Twips(680)
        section.header_distance, section.footer_distance = Twips(100), Twips(113)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(9)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def paragraph(text='', bold=False, color='000000', size=9, center=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = bold, Pt(size), RGBColor.from_string(color)
        return p

    def heading(number, title, new_page=False):
        p = paragraph(f'{number}.     {title}', bold=True, size=10)
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.page_break_before = new_page

    def fill(cell, content='', bold=False, bg=None, color='000000', center=False):
        cell.text = ''
        set_cell_margins(cell, top=40, bottom=40, left=75, right=75)
        if bg:
            set_cell_background(cell, bg)
        content = _text(content)
        if '<' in content and '>' in content:
            append_html_to_cell_or_paragraph(cell, content, default_font_size=9)
            return
        for i, line in enumerate(content.split('\n') or ['']):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            append_mixed_text(p,line,font_name='Arial',font_size=9,bold=bold,font_color=RGBColor.from_string(color))

    def bullets(cell, items, colors=False):
        fill(cell)
        for i, item in enumerate(items):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            append_mixed_text(p,'• ' + _text(item),font_name='Arial',font_size=9,
                              font_color=RGBColor.from_string(COLORS[i % 4] if colors else '000000'))

    def table(rows, weights, headers=None):
        t = doc.add_table(rows=rows, cols=len(weights))
        t.autofit = False
        widths = [TABLE_WIDTH * w // sum(weights) for w in weights]
        widths[-1] += TABLE_WIDTH - sum(widths)
        set_table_col_widths_and_indent(t, widths, indent_twip=0)
        add_table_borders_black(t, sz='4')
        for row in t.rows:
            for cell in row.cells:
                fill(cell)
        if headers:
            repeat = OxmlElement('w:tblHeader')
            t.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell, title in zip(t.rows[0].cells, headers):
                fill(cell, title, bold=True, bg='BDD7EE', center=True)
        return t

    admin, ident = container.administrativeContext, container.identity
    types = {'unit': 'UNIDAD DIDÁCTICA DE APRENDIZAJE', 'project': 'PROYECTO DE APRENDIZAJE',
             'learning_experience': 'EXPERIENCIA DE APRENDIZAJE', 'context': 'PLANIFICACIÓN CONTEXTUALIZADA'}
    p = paragraph(types[ident.planningType], bold=True, size=12, center=True)
    p.paragraph_format.keep_with_next = True
    p = paragraph(ident.title or 'Título por completar', bold=True, color='0070C0', size=12, center=True)
    p.runs[0].italic = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(7)
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'EAF6FC')
    p._p.get_or_add_pPr().append(shade)

    # Same three-row merged information grid as the supplied reference.
    info = table(3, [2170,3694,72,816,99,246,1485,1339,1136,1083,3504])
    areas = ', '.join(dict.fromkeys(e.area.officialName for e in container.curriculumMap))
    duration = ident.duration or {}
    dates = ' al '.join(filter(None, (ident.startDate, ident.endDate)))
    duration_label = dates or f"{duration.get('value', 0)} " + {'weeks':'semanas','days':'días','sessions':'sesiones'}.get(duration.get('unit'),'')
    level = {'secondary':'SECUNDARIA','primary':'PRIMARIA','initial':'INICIAL'}.get(ident.level,ident.level)
    rows = [
        [(0,0,'Institución Educativa',True),(1,5,admin.institution,False),(6,6,'Director',True),(7,10,admin.director,False)],
        [(0,0,'Coordinador',True),(1,1,admin.coordinator,False),(2,3,'Docente',True),(4,8,admin.teacher,False),(9,9,'Nivel',True),(10,10,level,False)],
        [(0,0,'Área(s)',True),(1,2,areas,False),(3,4,'Grado',True),(5,6,f'{ident.grade or ""}.º',False),(7,7,'Sección(es)',True),(8,8,', '.join(admin.sections),False),(9,9,'Duración',True),(10,10,duration_label,False)]]
    for ri, fields in enumerate(rows):
        for start,end,text,label in fields:
            cell = info.cell(ri,start)
            if end > start:
                cell = cell.merge(info.cell(ri,end))
            fill(cell,text,bold=label,bg='E2F0D9' if label else 'E2F5F1')
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)

    heading('I','SITUACIÓN SIGNIFICATIVA')
    sit = container.significantSituation
    parts = dict.fromkeys(filter(None,[sit.context,sit.problemOrOpportunity,sit.relevance,sit.studentRole,sit.expectedResponse]))
    html = ''.join(f'<p>{escape(item)}</p>' for item in parts)
    if container.drivingQuestion:
        html += f'<p><strong>{escape(container.drivingQuestion)}</strong></p>'
    fill(table(1,[1]).cell(0,0),html)

    heading('II','PROPÓSITO Y PRODUCTO FINAL DE LA UNIDAD DIDÁCTICA')
    t = table(2,[45,55],['PROPÓSITO DE LA EXPERIENCIA DE APRENDIZAJE','PRODUCTO DE LA EXP. DE APREND.'])
    fill(t.cell(1,0),'\n'.join(dict.fromkeys(filter(None,[container.purpose.summary,container.purpose.what,container.purpose.why,container.purpose.context]))))
    product = container.finalProduct
    if product:
        fill(t.cell(1,1),[product.title,product.description,product.expectedComponents])

    heading('III','ESTÁNDAR DE APRENDIZAJE',True)
    curriculum = container.curriculumMap
    t = table(1+max(1,len(curriculum)),[11,28,61],['COMPETENCIAS','CAPACIDADES','ESTÁNDARES DE APRENDIZAJE DE LAS COMPETENCIAS'])
    for i, entry in enumerate(curriculum,1):
        fill(t.cell(i,0),entry.competency.officialName,True,'BDD7EE',center=True)
        bullets(t.cell(i,1),[c.officialName for c in entry.capacities],True)
        fill(t.cell(i,2),entry.standard.description if entry.standard else '')

    heading('IV','COMPETENCIAS Y ENFOQUES TRANSVERSALES')
    t = table(8,[12,19,69],['COMPETENCIAS TRANSVERSALES','CAPACIDADES','DESEMPEÑOS'])
    transverse = [
        ('Se desenvuelve en entornos virtuales generados por las TIC',['Personaliza entornos virtuales','Gestiona información del entorno virtual','Interactúa en entornos virtuales','Crea objetos virtuales en diversos formatos']),
        ('Gestiona su aprendizaje de manera autónoma',['Define metas de aprendizaje','Organiza acciones estratégicas para alcanzar sus metas de aprendizaje','Monitorea y ajusta su desempeño durante el proceso de aprendizaje'])]
    default_performances = [
        ['Clasifica y organiza información de diversas fuentes reconociendo derechos de autor y citando fuentes confiables.',
         'Registra datos mediante hojas de cálculo que le permitan ordenar y secuenciar información relevante.',
         'Utiliza herramientas multimedia e interactivas para elaborar proyectos escolares y representaciones digitales.'],
        ['Determina metas de aprendizaje viables asociadas a sus conocimientos y habilidades formulándose preguntas de manera reflexiva.',
         'Organiza un conjunto de estrategias y procedimientos en función del tiempo y de los recursos de que dispone.',
         'Explica las acciones realizadas y los recursos movilizados en función de su pertinencia al logro de las metas de aprendizaje.']]
    start = 1
    for i,(name,capacities) in enumerate(transverse):
        end = start + len(capacities) - 1
        fill(t.cell(start,0).merge(t.cell(end,0)),name,True,center=True)
        for offset,capacity in enumerate(capacities):
            fill(t.cell(start+offset,1),capacity,color=COLORS[offset % 4],center=True)
        matches = [e.valuesOrAttitudes for e in container.transversalElements if e.name.casefold()==name.casefold()]
        bullets(t.cell(start,2).merge(t.cell(end,2)),matches or default_performances[i],True)
        start = end + 1
    paragraph()
    approaches = [e for e in container.transversalElements if e.name.casefold() not in {x[0].casefold() for x in transverse}]
    t = table(1+max(1,sum(max(1,len(e.valuesOrAttitudes)) for e in approaches)),[12,23,65],['ENFOQUE','VALORES','PRÁCTICA DE VALORES'])
    start = 1
    for item in approaches:
        values = item.valuesOrAttitudes or ['']
        end = start + len(values) - 1
        name_cell = t.cell(start,0).merge(t.cell(end,0)) if end > start else t.cell(start,0)
        fill(name_cell,item.name,True,center=True)
        for offset,value in enumerate(values):
            fill(t.cell(start+offset,1),value)
            practices = item.practices[offset:offset+1] if len(item.practices)==len(values) else item.practices
            fill(t.cell(start+offset,2),practices or 'Demuestra compromiso y cooperación en las actividades pedagógicas individuales y grupales.')
        start = end + 1

    heading('V','SECUENCIA DE LAS SESIONES DE APRENDIZAJE',True)
    sequence = sorted(container.sequence,key=lambda item:item.index)
    t = table(1+max(1,len(sequence)),[12,13,12,28,13,12,10],['Competencia','Semana/fecha','Nombre de la sesión','Criterios','Campo temático','Evidencia','Instrumento de evaluación'])
    by_map = {e.id:e for e in curriculum}
    for ri,item in enumerate(sequence,1):
        entries = [by_map[ref] for ref in item.curriculumMapRefs if ref in by_map]
        if not entries:
            entries = [e for e in curriculum if e.competency.id in item.competencyRefs]
        fill(t.cell(ri,0),[e.competency.officialName for e in entries],True,center=True)
        fill(t.cell(ri,1),[f'SEMANA {item.week}' if item.week else '',item.date],True,color='0070C0',center=True)
        fill(t.cell(ri,2),f'Sesión de aprendizaje {item.index:02}:\n\n{item.title}',True,center=True)
        criteria = [c for e in entries for c in e.criteria if c.id in item.criterionRefs]
        bullets(t.cell(ri,3),[c.description for c in criteria],True)
        bullets(t.cell(ri,4),item.knowledge)
        evidence = [_text(e) for e in item.evidence] + [_text(item.partialProduct)]
        fill(t.cell(ri,5),'\n'.join(dict.fromkeys(filter(None,evidence))),center=True)
        instruments = [_text(i) for i in item.assessmentInstruments]
        fill(t.cell(ri,6),instruments if any(instruments) else [c.instrument for c in criteria if c.instrument],center=True)

    heading('VI','MATERIALES, RECURSOS EDUCATIVOS Y ORIENTACIONES PARA LA EVALUACIÓN')
    t = table(3,[44,56],['MATERIALES Y RECURSOS EDUCATIVOS','ORIENTACIONES PARA LA EVALUACIÓN'])
    bullets(t.cell(1,0),[_text(i.model_dump() if hasattr(i,'model_dump') else i) for i in container.resources])
    assessment = container.assessmentPlan
    fill(t.cell(1,1),[assessment.formativeAssessment,assessment.feedbackApproach,assessment.selfAssessment,assessment.peerAssessment,assessment.teacherAssessment,assessment.recommendedInstruments])
    bibliography = t.cell(2,0).merge(t.cell(2,1))
    fill(bibliography,'BIBLIOGRAFÍA\n'+_text([i.model_dump() if hasattr(i,'model_dump') else i for i in container.bibliography]))
    paragraph()
    p = paragraph('_______________________________\t\t_______________________________',center=True)
    p.paragraph_format.keep_with_next = True
    paragraph(f'{admin.coordinator or "Coordinador(a)"}\t\t{admin.teacher or "Docente"}',center=True)
    paragraph('Coordinador(a) pedagógico(a)\t\tDocente',center=True)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
